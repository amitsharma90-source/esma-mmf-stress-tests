"""Scenario audit workbook - the working behind the macro scenarios (MST-01, MST-02).

Two sheets, one per natural grain:

    mst_summary     one row per fund      the market shock build-up, the liquidity cost,
                                           the pre- and post-shock weekly liquid asset
                                           buckets, and both MST-01/MST-02 results
    mst_positions   one row per holding   each position's own rate, credit and FX legs,
                                           its stressed market value, and the WLA bucket it
                                           lands in before and after the shock

Every figure comes from the engine's own functions - `weekly_liquid_assets`, and the
market-shock, adverse-FX and stressed-value functions `scenarios.run` itself calls. Nothing
here can change an engine result; this only exposes the arithmetic paras 63 and 65 combine
into two numbers.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .classify import weekly_liquid_assets, wla_bucket
from .loader import Fund
from .measures import PositionMeasures, blended_redemption
from .reference import Reference
from .scenarios import FX_ATTR, adverse_fx, fx_pnl, impact, market_shock, stressed_values

NAVY = "FF0E2C4B"
AMBER = "FFFFF2CC"
WHITE_BOLD = Font(color="FFFFFFFF", bold=True, size=10)


def _sheet(wb, name, columns, note):
    """columns: sequence of (header, width, number_format or None)."""
    ws = wb.create_sheet(name)
    ws.append([note])
    ws["A1"].font = Font(italic=True, size=9, color="FF555555")
    ws.append([])
    ws.append([c for c, _w, _f in columns])
    for i in range(1, len(columns) + 1):
        c = ws.cell(row=3, column=i)
        c.font = WHITE_BOLD
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="left", vertical="center")
    for i, (_c, w, _f) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "D4"
    return ws


def _apply_formats(ws, columns):
    """Formats go only on cells holding a value - an empty typed cell voids the file."""
    for row in range(4, ws.max_row + 1):
        for i, (_c, _w, fmt) in enumerate(columns, start=1):
            if fmt and ws.cell(row=row, column=i).value is not None:
                ws.cell(row=row, column=i).number_format = fmt


def _shade(ws, ncols, col, value, colour):
    """Shade a whole row. See audit.py for why empty cells are typed as string first."""
    fill = PatternFill("solid", fgColor=colour)
    for row in range(4, ws.max_row + 1):
        if ws.cell(row=row, column=col).value == value:
            for c in range(1, ncols + 1):
                cell = ws.cell(row=row, column=c)
                cell.fill = fill
                if cell.value is None:
                    cell.data_type = "s"


def write(path: Path, funds: list[Fund], ref: Reference,
          measures: dict[str, dict[str, PositionMeasures]]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    w1 = ref.bucket_weight(1)

    def sum_(fcode, attr):
        return sum(getattr(m, attr) or 0.0 for m in measures[fcode].values())

    # ---- 1. one row per fund ------------------------------------------------
    cols = [
        ("fund", 14, None), ("NAV", 18, "#,##0.00"),
        ("interest_ist01_total", 20, "#,##0.00"), ("credit_cst01_total", 19, "#,##0.00"),
        ("fx_appreciation_total", 21, "#,##0.00"), ("fx_depreciation_total", 21, "#,##0.00"),
        ("fx_leg_used", 15, None), ("fx_adverse", 16, "#,##0.00"),
        ("market_shock", 18, "#,##0.00"),
        ("macro_redemption", 18, "0.0%"),
        ("liquidity_lst01_mst_total", 25, "#,##0.00"), ("liquidity_cost", 17, "#,##0.00"),
        ("MST01_impact_on_NAV", 20, "0.00%"),
        ("stressed_NAV", 16, "#,##0.00"),
        ("bucket_1_raw", 16, "#,##0.00"), ("bucket_2_raw", 16, "#,##0.00"),
        ("bucket_1_stressed", 18, "#,##0.00"), ("bucket_2_stressed", 18, "#,##0.00"),
        ("MST02_outflow", 16, "#,##0.00"),
        ("MST02_first", 13, "0.00%"), ("MST02_total", 13, "0.00%"),
    ]
    ws = _sheet(wb, "mst_summary", cols,
                "One row per fund. Amounts are P&L - a loss is negative. para 63/65: the market shock "
                "is rates + credit + the ADVERSE side of FX (whichever of Table 10 appreciation and "
                "Table 11 depreciation costs the fund more); SST is deliberately excluded, because on "
                "fixed-rate paper it equals IST-01 exactly and would double count the same curve "
                "move. liquidity_cost is the macro redemption rate applied to the fund's "
                "liquidity_lst01_mst (a vertical slice, not the weekly-tradable subset RST-01 uses). "
                "MST-01 = -(market_shock + liquidity_cost) / NAV, so a loss reads positive (para 58). "
                "MST-02 revalues every position by its own rate, credit and adverse FX legs, "
                "re-sums the WLA buckets on those stressed values, and divides by the macro outflow "
                "measured against the STRESSED NAV (both sides of the ratio are post-shock, per "
                "para 63).")
    for f in funds:
        nav = f.nav
        m_f = measures[f.code]
        fx_leg, fx_adverse = adverse_fx(f, m_f)
        shock = market_shock(f, m_f)
        macro_redemption = blended_redemption(f, ref, macro=True)
        liq_mst_total = sum_(f.code, "liquidity_lst01_mst")
        liquidity_cost = macro_redemption * liq_mst_total
        stressed_nav = nav + shock
        wla_raw = weekly_liquid_assets(f, ref)
        wla_stressed = weekly_liquid_assets(f, ref, stressed_values(f, m_f))
        outflow = macro_redemption * stressed_nav
        ws.append([
            f.code, nav, sum_(f.code, "interest_rates_ist01"), sum_(f.code, "credit_cst01"),
            fx_pnl(f, m_f, "APPRECIATION"), fx_pnl(f, m_f, "DEPRECIATION"), fx_leg, fx_adverse,
            shock, macro_redemption, liq_mst_total, liquidity_cost,
            impact(shock + liquidity_cost, nav), stressed_nav,
            wla_raw["bucket_1"], wla_raw["bucket_2"],
            wla_stressed["bucket_1"], wla_stressed["bucket_2"],
            outflow,
            wla_stressed["bucket_1"] * w1 / outflow if outflow > 0 else 0.0,
            wla_stressed["total"] / outflow if outflow > 0 else 0.0,
        ])
    _apply_formats(ws, cols)

    # ---- 2. one row per position ---------------------------------------------
    cols = [
        ("fund", 14, None), ("security_id", 13, None), ("instrument", 24, None),
        ("issuer", 38, None), ("security_currency", 17, None), ("market_value", 18, "#,##0.00"),
        ("interest_rates_ist01", 20, "#,##0.00"), ("credit_cst01", 16, "#,##0.00"),
        ("fx_fst01_value_change", 21, "0.000000"), ("fx_fst01_contribution", 21, "#,##0.00"),
        ("fx_fst02_value_change", 21, "0.000000"), ("fx_fst02_contribution", 21, "#,##0.00"),
        ("fx_adverse_contribution", 23, "#,##0.00"),
        ("stressed_market_value", 21, "#,##0.00"),
        ("liquidity_lst01_mst", 19, "#,##0.00"),
        ("wla_bucket", 12, None), ("wla_weight", 11, "0%"),
        ("wla_contribution_raw", 20, "#,##0.00"), ("wla_contribution_stressed", 25, "#,##0.00"),
    ]
    ws = _sheet(wb, "mst_positions", cols,
                "Every holding once. interest_rates_ist01 and credit_cst01 are this position's own "
                "contribution to the market shock (para 63). fx_fst01/02_value_change is the relative "
                "change in the holding's value in the fund's currency under Table 10 / Table 11 - a "
                "gain is positive, and a holding in the fund's own currency has none. "
                "fx_adverse_contribution applies the fund-level adverse leg, one direction for every "
                "holding. stressed_market_value is market_value plus the rate, credit and adverse FX "
                "legs, which is what MST-02's weekly liquid asset buckets are re-summed on. "
                "liquidity_lst01_mst is this position's contribution to the pool the macro redemption "
                "rate is applied to - it is not itself scaled by that rate. wla_bucket/weight do not "
                "change under the shock; only the amount they are weighted against does.")
    for f in funds:
        m_f = measures[f.code]
        leg, _fx = adverse_fx(f, m_f)
        stressed = stressed_values(f, m_f)
        for p in f.positions:
            m = m_f[p.security_id]
            b = wla_bucket(p, ref)
            ws.append([
                f.code, p.security_id, p.instrument, p.issuer, p.currency, p.market_value,
                m.interest_rates_ist01, m.credit_cst01,
                m.fx_fst01, m.fx_fst01 * p.market_value,
                m.fx_fst02, m.fx_fst02 * p.market_value,
                getattr(m, FX_ATTR[leg]) * p.market_value,
                stressed[p.security_id], m.liquidity_lst01_mst,
                b.bucket, b.weight, p.market_value * b.weight,
                stressed[p.security_id] * b.weight,
            ])
    _apply_formats(ws, cols)
    _shade(ws, len(cols), 16, "None", AMBER)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
