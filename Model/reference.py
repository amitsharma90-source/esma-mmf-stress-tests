"""Load esma_reference.xlsx and answer every calibration question against it.

This module is the only place the engine learns a number, and it holds none of its own.
Each lookup states which table it read and why, so a result can be traced back to a
paragraph of the Guidelines or an article of the Regulation.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


class ReferenceError(RuntimeError):
    """A lookup found no row. Never silently returns zero - a missing shock is a bug."""


def _sheet(wb, name: str, first_col: str) -> list[dict]:
    """Read a sheet into dicts, skipping the italic note row the builder writes on top."""
    if name not in wb.sheetnames:
        raise ReferenceError(f"esma_reference.xlsx has no sheet {name!r}")
    rows = list(wb[name].iter_rows(values_only=True))
    try:
        h = next(i for i, r in enumerate(rows) if r and r[0] == first_col)
    except StopIteration:
        raise ReferenceError(f"{name}: no header row starting with {first_col!r}") from None
    header = [c for c in rows[h] if c is not None]
    out = []
    for r in rows[h + 1:]:
        if r is None or r[0] is None:
            continue
        out.append({header[i]: r[i] for i in range(len(header))})
    return out


@dataclass(frozen=True)
class Shock:
    """A looked-up value plus the table it came from, so results stay auditable."""
    value: float | None
    source: str


class Reference:
    def __init__(self, path: str | Path):
        wb = load_workbook(Path(path), read_only=True, data_only=True)
        self.path = Path(path)

        self.ldf_country = {r["country"]: r for r in _sheet(wb, "table_01_sovereign_ldf_country", "country")}
        self.ldf_sov_rating = {r["lt_rating"]: r for r in _sheet(wb, "table_02_sovereign_ldf_rating", "lt_rating")}
        self.ldf_corp_rating = {r["lt_rating"]: r for r in _sheet(wb, "table_03_corporate_ldf_rating", "lt_rating")}
        self.pip = {r["asset_category"]: r["price_impact_parameter"]
                    for r in _sheet(wb, "table_04_price_impact", "asset_category")}
        self.spread_country = {r["country"]: r for r in _sheet(wb, "table_05_govt_spread_country", "country")}
        self.spread_rating = {r["lt_rating"]: r for r in _sheet(wb, "table_06_corp_spread_rating", "lt_rating")}
        self.lgd = {r["seniority"]: r["loss_given_default"] for r in _sheet(wb, "table_07_lgd", "seniority")}
        self.rate_ccy = {r["currency"]: r for r in _sheet(wb, "table_08_rate_shock_currency", "currency")}
        self.rate_default = {r["area"]: r for r in _sheet(wb, "table_09_rate_shock_default", "area")}
        self.fx_appreciation = {r["pair"]: r["shock"] for r in _sheet(wb, "table_10_fx_eur_appreciation", "pair")}
        self.fx_depreciation = {r["pair"]: r["shock"] for r in _sheet(wb, "table_11_fx_eur_depreciation", "pair")}
        self.redemption = {r["investor_type"]: r["net_outflow"]
                           for r in _sheet(wb, "table_13_redemption_standalone", "investor_type")}
        self.redemption_macro = {r["investor_type"]: r["net_outflow"]
                                 for r in _sheet(wb, "table_14_redemption_macro", "investor_type")}

        self.wla_rules = _sheet(wb, "wla_bucket_rules", "bucket")
        self.article_limits = {(r["fund_type"], r["nav_type"]): r
                               for r in _sheet(wb, "article_limits", "fund_type")}
        self.tenor_maps = {r["maturity_bucket"]: r for r in _sheet(wb, "tenor_maps", "maturity_bucket")}
        self.exclusions = {r["instrument_type"]: r for r in _sheet(wb, "exclusion_matrix", "instrument_type")}
        self.sensitivity = {r["rate_type"]: r for r in _sheet(wb, "sensitivity_basis", "rate_type")}
        self.settings = {r["setting"]: r["value"] for r in _sheet(wb, "engine_settings", "setting")}
        self.article_17_7 = {r["issuer_type"]: bool(r["qualifies"])
                             for r in _sheet(wb, "article_17_7_issuer_types", "issuer_type")}
        self.country_codes = {r["country"]: r["iso_alpha_2"] for r in _sheet(wb, "country_codes", "country")}
        self.liquidity_bucket_days = {r["liquidity_bucket"]: int(r["max_days"])
                                      for r in _sheet(wb, "liquidity_buckets", "liquidity_bucket")}
        self.ldf_mapping = {r["issuer_type"]: r["source"] for r in _sheet(wb, "ldf_mapping", "issuer_type")}
        self._cqs = _sheet(wb, "cqs_mapping", "cqs")
        self._credit_map = _sheet(wb, "credit_spread_mapping", "issuer_type")
        self._pip_map = _sheet(wb, "price_impact_mapping", "instrument_type")
        wb.close()

        # Art 10 money market instrument eligibility, read off article_limits rather than
        # restated: 397 days for paper that does not reset (the short-term instrument limit,
        # which binds every fund), and two years for paper resetting within 397 days.
        rows = list(self.article_limits.values())
        resets = [int(r["max_reset_days"]) for r in rows if r["max_reset_days"] is not None]
        if not rows or not resets:
            raise ReferenceError("article_limits must carry instrument maturity and reset limits")
        self.mmi_max_residual_days = min(int(r["max_instrument_maturity_days"]) for r in rows)
        self.mmi_max_extended_days = max(int(r["max_instrument_maturity_days"]) for r in rows)
        self.mmi_max_reset_days = min(resets)

    def bucket_weight(self, bucket: int) -> float:
        """para 61's weight for a bucket - x100% or x85% - taken from its rule lines."""
        weights = {float(r["weight"]) for r in self.wla_rules if int(r["bucket"]) == bucket}
        if len(weights) != 1:
            raise ReferenceError(f"wla_bucket_rules bucket {bucket}: expected one weight, "
                                 f"found {sorted(weights)}")
        return next(iter(weights))

    # -- scope ------------------------------------------------------------
    def in_scope(self, instrument: str, risk: str) -> bool:
        row = self.exclusions.get(instrument)
        if row is None:
            raise ReferenceError(f"exclusion_matrix has no row for instrument {instrument!r}")
        return bool(row[f"in_scope_{risk}"])

    def is_maturing(self, instrument: str) -> bool:
        return bool(self.exclusions[instrument]["is_maturing"])

    def diversification_scope(self, instrument: str) -> str:
        """Art 17(1)(a), which groups MMIs, securitisations and ABCPs together."""
        return self.exclusions[instrument]["diversification_scope"]

    def asset_category(self, instrument: str) -> str:
        """Art 9(1), where those are SEPARATE categories - (a) money market instruments,
        (b) eligible securitisations and ABCPs. para 61 bucket 2 gives them different
        lines with different CQS tests, so they must not be conflated here."""
        return self.exclusions[instrument]["mmfr_asset_category"]

    def qualifies_article_17_7(self, instrument: str, issuer_type: str) -> bool:
        return (bool(self.exclusions[instrument]["article_17_7_eligible"])
                and self.article_17_7.get(issuer_type, False))

    def tenor(self, maturity_bucket: str, family: str) -> str:
        row = self.tenor_maps.get(maturity_bucket)
        if row is None:
            raise ReferenceError(f"tenor_maps has no row for bucket {maturity_bucket!r}")
        return row[f"{family}_tenor"]

    # -- lookups ----------------------------------------------------------
    def cqs(self, st_rating: str | None) -> int | None:
        """Credit Quality Step from the short-term rating. Only the WLA buckets use it."""
        if not st_rating:
            return None
        for r in self._cqs:
            if st_rating in (r["moodys_st"], r["sp_st"], r["fitch_st"]):
                return int(r["cqs"])
        return None

    def liquidity_discount(self, instrument, issuer_type, country, lt_rating,
                           maturity_bucket) -> Shock:
        if not self.in_scope(instrument, "ldf"):
            return Shock(None, f"out of liquidity scope ({instrument})")
        source = self.ldf_mapping.get(issuer_type, "TABLE_03_RATING")
        tenor = self.tenor(maturity_bucket, "ldf")

        if source == "TABLE_01_COUNTRY_ELSE_02_RATING":
            code = self.country_codes.get(country)
            if code and code in self.ldf_country:
                return Shock(self.ldf_country[code][tenor], f"Table 1 {code} {tenor}")
            source = "TABLE_02_RATING"
        if source == "TABLE_02_RATING":
            row = self.ldf_sov_rating.get(lt_rating)
            if row is None:
                raise ReferenceError(f"Table 2 has no rating {lt_rating!r}")
            return Shock(row[tenor], f"Table 2 {lt_rating} {tenor}")
        if source == "TABLE_03_RATING":
            row = self.ldf_corp_rating.get(lt_rating)
            if row is None:
                raise ReferenceError(f"Table 3 has no rating {lt_rating!r}")
            return Shock(row[tenor], f"Table 3 {lt_rating} {tenor}")
        return Shock(None, "no liquidity discount table applies")

    def price_impact(self, instrument: str, sector: str | None) -> Shock:
        if not self.in_scope(instrument, "mif"):
            return Shock(None, f"out of price impact scope ({instrument})")
        for r in self._pip_map:
            if r["instrument_type"] == instrument and r["industry_sector"] in ("*", sector):
                cat = r["table_04_category"]
                return Shock(self.pip.get(cat), f"Table 4 {cat}")
        raise ReferenceError(f"price_impact_mapping has no row for {instrument!r}/{sector!r}")

    def credit_spread(self, instrument, issuer_type, sector, country, lt_rating,
                      maturity_bucket) -> Shock:
        if not self.in_scope(instrument, "credit"):
            return Shock(None, f"out of credit scope ({instrument})")
        for r in self._credit_map:
            if r["issuer_type"] != issuer_type or r["industry_sector"] not in ("*", sector):
                continue
            if r["source"] == "TABLE_05_COUNTRY":
                tenor = self.tenor(maturity_bucket, "credit_spread")
                row = self.spread_country.get(country)
                if row is None:
                    # A country ESMA does not list falls back to its weighted-average area.
                    row = self.spread_country.get("EU (weighted averages)")
                    if row is None:
                        raise ReferenceError(f"Table 5 has no row for {country!r} and no EU fallback")
                    return Shock(row[tenor], f"Table 5 EU average {tenor} (fallback for {country})")
                return Shock(row[tenor], f"Table 5 {country} {tenor}")
            if r["source"] == "TABLE_06_RATING":
                col = r["table_06_column"]
                row = self.spread_rating.get(lt_rating)
                if row is None:
                    raise ReferenceError(f"Table 6 has no rating {lt_rating!r}")
                return Shock(row[col], f"Table 6 {lt_rating} {col}")
            return Shock(None, "no credit spread table applies")
        raise ReferenceError(f"credit_spread_mapping has no row for {issuer_type!r}/{sector!r}")

    def rate_shock(self, instrument: str, currency: str, maturity_bucket: str) -> Shock:
        if not self.in_scope(instrument, "rates"):
            return Shock(None, f"out of interest rate scope ({instrument})")
        tenor = self.tenor(maturity_bucket, "rate_shock")
        row = self.rate_ccy.get(currency)
        if row is not None:
            return Shock(row[tenor], f"Table 8 {currency} {tenor}")
        row = self.rate_default.get("Other advanced economies")
        if row is None:
            raise ReferenceError(f"no Table 8 row for {currency!r} and no Table 9 default")
        return Shock(row[tenor], f"Table 9 advanced-economy default {tenor} (no {currency} row)")

    def loss_given_default(self, instrument: str, seniority: str) -> Shock:
        if not self.in_scope(instrument, "credit"):
            return Shock(None, f"out of credit scope ({instrument})")
        if seniority not in self.lgd:
            raise ReferenceError(f"Table 7 has no seniority {seniority!r}")
        return Shock(self.lgd[seniority], f"Table 7 {seniority}")

    def fx_value_change(self, instrument: str, currency: str, base_currency: str,
                        direction: str) -> Shock:
        """Relative change in a holding's value, measured in the fund's currency, under
        Table 10 (EUR appreciation, FST-01) or Table 11 (EUR depreciation, FST-02).
        A gain is positive.

        Each pair is quoted in market convention and shocked as a relative move in the rate:
            EURUSD  1 EUR = x USD   the euro's value in dollars moves by the shock
            USDJPY  1 USD = x JPY   x rising means each yen is worth less: 1/(1+s) - 1
            EURGBP  1 EUR = x GBP   no dollar quote, so sterling is reached through the euro
        Starting from the fund's currency at no change, the pairs are walked one link at a
        time until the holding's currency is reached. For a pair XY moving by s, X's value
        in Y moves by s, so a known Y gives X = (1+s)(1+Y) - 1 and a known X gives
        Y = (1+X)/(1+s) - 1. A currency no chain reaches is an error, never a zero.
        """
        if not self.in_scope(instrument, "fx"):
            return Shock(None, f"out of FX scope ({instrument})")
        if currency == base_currency:
            return Shock(0.0, f"no currency exposure ({currency})")
        table, name = ((self.fx_appreciation, "Table 10") if direction == "APPRECIATION"
                       else (self.fx_depreciation, "Table 11"))

        change = {base_currency: 0.0}
        route: dict[str, list[str]] = {base_currency: []}
        frontier = [base_currency]
        while frontier and currency not in change:
            reached = []
            for known in frontier:
                for pair, s in table.items():
                    x, y = pair[:3], pair[3:]
                    if y == known and x not in change:
                        new, value = x, (1.0 + s) * (1.0 + change[y]) - 1.0
                    elif x == known and y not in change:
                        new, value = y, (1.0 + change[x]) / (1.0 + s) - 1.0
                    else:
                        continue
                    change[new] = value
                    route[new] = route[known] + [f"{pair} {s:+.0%}"]
                    reached.append(new)
            frontier = reached
        if currency not in change:
            raise ReferenceError(f"{name} has no pair linking {currency} to {base_currency}")
        return Shock(change[currency], f"{name} {currency} in {base_currency} via "
                                       + ", ".join(route[currency]))

    def limits(self, fund_type: str, nav_type: str) -> dict:
        key = (fund_type, nav_type)
        if key not in self.article_limits:
            raise ReferenceError(f"article_limits has no row for {key}")
        return self.article_limits[key]
