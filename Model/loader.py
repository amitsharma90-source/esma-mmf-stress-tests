"""Read the position file and the fund static file, and derive what the input omits.

Three derivations matter:

  residual maturity  legal maturity minus the valuation date, on ACT/365
  rate duration      to the NEXT RESET for floating paper, to legal maturity for fixed
  spread duration    always to legal maturity

The second and third are why the output carries two sensitivity columns. MMFR Art 2(19)
defines WAM on the next interest-rate reset when that is sooner; Art 2(20) defines WAL on
legal maturity alone. The same split separates rate risk from spread risk: a floating
coupon resets away a curve move but keeps its benchmark spread to maturity.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from .reference import Reference

DAYS_PER_YEAR = 365.0

# engine_settings.working_day_calendar -> numpy weekmask, Monday first.
WEEKMASKS = {"WEEKENDS_ONLY": "1111100"}

# security_type in the file -> instrument_type in exclusion_matrix
INSTRUMENT_BY_SECURITY_TYPE = {
    "TBILL": "TREASURY_BILL", "GOVT": "GOVERNMENT_BOND", "SUPRA": "SUPRANATIONAL_BOND",
    "CD": "CERTIFICATE_OF_DEPOSIT", "CP": "COMMERCIAL_PAPER", "CORP": "CORPORATE_BOND",
    "FRN": "FLOATING_RATE_NOTE", "ABS": "SECURITISATION", "ABCP": "ABCP",
    "DEPO": "TIME_DEPOSIT", "REVREPO": "REVERSE_REPO", "MMFSHARE": "MMF_SHARES",
    "CASH": "CASH",
}


class LoaderError(RuntimeError):
    pass


@dataclass
class Position:
    fund: str
    fund_name: str
    security_id: str
    description: str
    issuer: str
    issuer_country: str
    issuer_type: str
    asset_class: str
    security_type: str
    instrument: str
    maturity_date: date
    rate_type: str
    next_reset_date: date | None
    reference_index: str | None
    maturity_bucket: str
    liquidity_bucket: str
    settlement_days: int
    termination_notice_days: int | None
    st_rating: str | None
    lt_rating: str | None
    sector: str | None
    seniority: str
    market_value: float
    currency: str                   # the holding's own currency, not the fund's
    yield_: float
    issuer_outstanding: float | None
    sts: bool = False               # Art 17(3): simple, transparent and standardised
    # derived
    residual_days: int = 0          # calendar days, for WAM/WAL, Art 10 and the 190-day tests
    maturity_working_days: int = 0  # working days, for the daily and weekly maturing tests
    reset_days: int | None = None
    rate_duration: float = 0.0
    spread_duration: float = 0.0
    cqs: int | None = None
    derived: dict = field(default_factory=dict)

    @property
    def wam_days(self) -> int:
        return self.reset_days if self.reset_days is not None else self.residual_days

    @property
    def wal_days(self) -> int:
        return self.residual_days


@dataclass
class Fund:
    code: str
    name: str
    fund_type: str
    nav_type: str
    mmfr_article: int
    base_currency: str
    institutional_share: float
    retail_share: float
    top2_investor_share: float
    nav_override: float | None
    positions: list[Position] = field(default_factory=list)

    @property
    def nav(self) -> float:
        """Sum of market values unless a real NAV was supplied.

        Keeping NAV equal to total assets is what lets the Article 24/25 percent-of-assets
        tests and the percent-of-NAV results share one denominator.
        """
        if self.nav_override:
            return float(self.nav_override)
        return sum(p.market_value for p in self.positions)

    @property
    def total_assets(self) -> float:
        return sum(p.market_value for p in self.positions)


def _as_date(v) -> date | None:
    if v is None:
        return None
    return v.date() if hasattr(v, "date") else v


def _rows(path: Path, sheet: str) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise LoaderError(f"{path.name} has no sheet {sheet!r}")
    rows = list(wb[sheet].iter_rows(values_only=True))
    wb.close()
    header = list(rows[0])
    return [{header[i]: r[i] for i in range(len(header))}
            for r in rows[1:] if r and r[0] is not None]


def working_days_between(start: date, end: date, calendar: str) -> int:
    """Working days from `start` to `end`: maturing tomorrow is 1, maturing today is 0.

    numpy counts working days in [start, end), so a Friday valuation reaching a Monday
    maturity counts only the Friday - one working day, which is how para 61 and Articles
    24/25 read. Calendar days would call the same Monday three days away.
    """
    if calendar not in WEEKMASKS:
        raise LoaderError(f"working_day_calendar {calendar!r} is not supported")
    if end <= start:
        return 0
    return int(np.busday_count(start, end, weekmask=WEEKMASKS[calendar]))


def _flag(value, what: str) -> bool:
    """A yes/no column. Blank means no - for STS that is the conservative reading."""
    if value is None or value == "":
        return False
    if isinstance(value, bool):
        return value
    text = str(value).strip().upper()
    if text in ("TRUE", "YES", "Y", "1"):
        return True
    if text in ("FALSE", "NO", "N", "0"):
        return False
    raise LoaderError(f"{what} = {value!r}, expected TRUE/FALSE or blank")


def resolve_fund_static(extracted: str | Path) -> Path:
    """The fund static workbook is hand-maintained, so its name is not ours to fix.

    Both spellings have been in use - the generated `fund_static.xlsx` and the
    hand-kept `Fund static.xlsx`. Accept either rather than failing the run on a
    rename; prefer the generated name when both are present.
    """
    d = Path(extracted)
    for name in ("fund_static.xlsx", "Fund static.xlsx"):
        cand = d / name
        if cand.is_file():
            return cand
    raise FileNotFoundError(
        f"no fund static workbook in {d} (looked for fund_static.xlsx and 'Fund static.xlsx')")


def load(position_path: str | Path, fund_static_path: str | Path,
         reference: Reference, valuation_date: date) -> list[Fund]:
    position_path, fund_static_path = Path(position_path), Path(fund_static_path)

    funds: dict[str, Fund] = {}
    for r in _rows(fund_static_path, "fund_static"):
        # Every share on this file is a PERCENTAGE, named with a _pct suffix so the scale
        # cannot be misread, and all three are divided by 100 in one place. They used to
        # disagree - two on 0-100 and one on 0-1 - which is a 100x error waiting to happen.
        shares = {k: r[k] for k in
                  ("institutional_share_pct", "retail_share_pct", "top2_investor_share_pct")}
        for k, v in shares.items():
            if v is None or not (0.0 <= float(v) <= 100.0):
                raise LoaderError(f"{r['fund_code']}: {k} = {v!r}, expected a percentage 0-100")
        share_sum = float(shares["institutional_share_pct"]) + float(shares["retail_share_pct"])
        if abs(share_sum - 100.0) > 1e-6:
            raise LoaderError(
                f"{r['fund_code']}: institutional + retail = {share_sum}, must be 100")
        funds[r["fund_code"]] = Fund(
            code=r["fund_code"], name=r["fund_name"], fund_type=r["fund_type"],
            nav_type=r["nav_type"], mmfr_article=int(r["mmfr_article"]),
            base_currency=r["base_currency"],
            institutional_share=float(shares["institutional_share_pct"]) / 100.0,
            retail_share=float(shares["retail_share_pct"]) / 100.0,
            top2_investor_share=float(shares["top2_investor_share_pct"]) / 100.0,
            nav_override=r["nav_override"],
        )

    calendar = reference.settings.get("working_day_calendar")
    for r in _rows(position_path, "positions"):
        code = r["portfolio_code"]
        if code not in funds:
            raise LoaderError(f"position {r['security_id']} references unknown fund {code!r}")

        # FX exposure is the gap between the currency a holding is denominated in and the
        # fund's own. Reading the holding's currency off the fund's column made every
        # holding match its fund, so FST-01/02 could never register anything.
        currency = r.get("security_currency_code")
        if not currency:
            raise LoaderError(f"{r['security_id']}: security_currency_code is required - the "
                              f"currency the holding is denominated in")
        fund_ccy = r.get("portfolio_currency_code")
        if fund_ccy and fund_ccy != funds[code].base_currency:
            raise LoaderError(f"{r['security_id']}: portfolio_currency_code {fund_ccy!r} does not "
                              f"match fund static base_currency {funds[code].base_currency!r}")

        sec_type = r["security_type"]
        if sec_type not in INSTRUMENT_BY_SECURITY_TYPE:
            raise LoaderError(f"unmapped security_type {sec_type!r} on {r['security_id']}")
        instrument = INSTRUMENT_BY_SECURITY_TYPE[sec_type]

        maturity = _as_date(r["maturity_date"])
        reset = _as_date(r["next_reset_date"])
        if maturity is None:
            raise LoaderError(f"{r['security_id']}: maturity_date is required")

        residual = max((maturity - valuation_date).days, 0)
        reset_days = max((reset - valuation_date).days, 0) if reset else None

        rate_type = r["rate_type"] or "FIXED"
        if rate_type == "FLOATING" and reset_days is None:
            raise LoaderError(f"{r['security_id']}: FLOATING requires next_reset_date")
        basis = reference.sensitivity.get(rate_type)
        if basis is None:
            raise LoaderError(f"sensitivity_basis has no row for rate_type {rate_type!r}")

        def days_for(which: str) -> int:
            return reset_days if basis[which] == "NEXT_RESET" and reset_days is not None else residual

        p = Position(
            fund=code, fund_name=r["portfolio_name"], security_id=r["security_id"],
            description=r["security_description"], issuer=r["issuer_name"],
            issuer_country=r["issuer_country"], issuer_type=r["issuer_type"],
            asset_class=r["asset_class"], security_type=sec_type, instrument=instrument,
            maturity_date=maturity, rate_type=rate_type, next_reset_date=reset,
            reference_index=r["reference_index"], maturity_bucket=r["maturity_bucket"],
            liquidity_bucket=r["liquidity_bucket"],
            settlement_days=int(r["settlement_days"]),
            termination_notice_days=(int(r["termination_notice_days"])
                                     if r["termination_notice_days"] is not None else None),
            st_rating=r["st_rating"], lt_rating=r["lt_rating"],
            sector=r["industry_sector"], seniority=r["seniority"] or "SENIOR",
            market_value=float(r["market_value"]), currency=str(currency).strip().upper(),
            yield_=float(r["yield"]),
            issuer_outstanding=(float(r["issuer_outstanding_amount"])
                                if r.get("issuer_outstanding_amount") else None),
            sts=_flag(r.get("sts"), f"{r['security_id']} sts"),
        )
        p.residual_days = residual
        p.maturity_working_days = working_days_between(valuation_date, maturity, calendar)
        p.reset_days = reset_days if rate_type == "FLOATING" else None
        p.rate_duration = days_for("rate_duration_to") / DAYS_PER_YEAR
        p.spread_duration = days_for("spread_duration_to") / DAYS_PER_YEAR
        p.cqs = reference.cqs(p.st_rating)
        funds[code].positions.append(p)

    empty = [c for c, f in funds.items() if not f.positions]
    if empty:
        raise LoaderError(f"funds with no positions: {empty}")
    return list(funds.values())
